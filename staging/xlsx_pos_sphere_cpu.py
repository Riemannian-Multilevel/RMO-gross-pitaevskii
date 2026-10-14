#!/usr/bin/env python3
"""Spreadsheet of the comparison tables 5-10 of tex/exact_line_search.tex (tab_cpu_rows_<metric>_<rule>.tex of
staging/tables_pos_sphere_cpu.py: mass, Frobenius and energy-adaptive, each with Armijo and the exact step), with the
same data sources and columns. Needs openpyxl.

Usage: staging/xlsx_pos_sphere_cpu.py <output.xlsx> <cpu_analysis.csv> <cpu_econn.csv> <cpu_exact.csv>

Sheets:
  all      one row per table cell (table, metric, rule, discretization, radius, retraction, version, status, reason,
           iterations, CPU time in ms, residual, energy, closed form, fastest positive), with filters;
  tables   the six tables in the layout of the LaTeX tables: CPU time in ms (until the stop for runs that did not
           converge), grey for L/Z, red for F, the fastest positive version in bold; the reason and residual are cell
           comments;
  legend   the meaning of the fields.
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tables_pos_sphere_cpu import DENSITY, DISCS, ECONN, NORM, PANEL, RADII, SHORT, load, stop_reason

TABLES = [(5, "mass", "armijo"), (6, "mass", "exact"), (7, "frobenius", "armijo"), (8, "frobenius", "exact"),
          (9, "ea", "armijo"), (10, "ea", "exact")]
METRIC_NAME = {"mass": "mass", "frobenius": "Frobenius", "ea": "energy-adaptive"}
RULE_NAME = {"armijo": "Armijo", "exact": "exact"}
FAMILY = {**{c: "density (m)" for c in DENSITY}, **{c: "exponential (e)" for c in ECONN},
          **{c: "normalization, positive" for c in NORM}, "unit": "normalization (unit sphere)"}
GREY = PatternFill("solid", fgColor="DDDDDD")
RED = PatternFill("solid", fgColor="F4B6B6")
BOLD = Font(bold=True)


def columns(sources, metric, rule):
    """(case, data, panel) as in tables_pos_sphere_cpu.rows()."""
    cols = []
    for c in ["unit"] + NORM + DENSITY + ECONN:
        for data, study in sources:
            p = PANEL[study].get((metric, rule))
            if p and any(c in data.get((p, d, r), {}) for d in DISCS for r in RADII):
                cols.append((c, data, p))
                break
    return cols


def status(r):
    if r["status"] == "converged":
        return "converged", ""
    if r["status"] == "failed":
        return "failed", "F"
    return "not converged", stop_reason(r)


def main():
    if len(sys.argv) != 5:
        sys.exit(__doc__)
    out = sys.argv[1]
    analysis, econn, exact = load(sys.argv[2]), load(sys.argv[3]), load(sys.argv[4])
    sources = [(analysis, "analysis"), (exact, "exact"), (econn, "econn")]

    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "all"
    header = ["table", "metric", "rule", "discretization", "radius", "retraction", "version", "status", "reason",
              "iterations", "CPU time [ms]", "residual", "energy", "closed form", "fastest positive"]
    ws_all.append(header)
    ws_tab = wb.create_sheet("tables")
    row = 1
    for number, metric, rule in TABLES:
        cols = columns(sources, metric, rule)
        title = f"Table {number}: {METRIC_NAME[metric]}, {RULE_NAME[rule]}"
        if rule == "exact":
            title += " (unit, norm-cap: closed form)"
        ws_tab.cell(row=row, column=1, value=title).font = BOLD
        row += 1
        heads = ["discretization", "radius"] + [SHORT[c] for c, _, _ in cols]
        for j, h in enumerate(heads, 1):
            ws_tab.cell(row=row, column=j, value=h).font = BOLD
        row += 1
        for d in DISCS:
            for rad in RADII:
                rs = [data.get((p, d, rad), {}).get(c) for c, data, p in cols]
                conv = [float(r["cpu_ms"]) if r and r["status"] == "converged" and c != "unit" else float("inf")
                        for r, (c, _, _) in zip(rs, cols)]
                best = min(conv)
                ws_tab.cell(row=row, column=1, value=d)
                ws_tab.cell(row=row, column=2, value=rad)
                for j, ((c, _, p), r) in enumerate(zip(cols, rs), 3):
                    if r is None:
                        continue
                    st, reason = status(r)
                    fastest = conv[j - 3] == best and best < float("inf")
                    closed = rule == "exact" and c in ["unit"] + NORM
                    ws_all.append([number, METRIC_NAME[metric], RULE_NAME[rule], d, rad, FAMILY[c], SHORT[c], st,
                                   reason, int(r["iterations"]), round(float(r["cpu_ms"]), 3),
                                   float(r["residual"]), float(r["energy"]), "yes" if closed else "no",
                                   "yes" if fastest else ""])
                    x = ws_tab.cell(row=row, column=j, value=None if st == "failed" else
                                    round(float(r["cpu_ms"]), 3))
                    x.number_format = "0.0##"
                    if st == "failed":
                        x.value = "F"
                        x.fill = RED
                        x.alignment = Alignment(horizontal="right")
                    elif st == "not converged":
                        x.fill = GREY
                        x.comment = Comment(f"{reason}: {'1000 iterations' if reason == 'L' else 'step 0 after '}"
                                            f"{'' if reason == 'L' else r['iterations'] + ' iterations'}; "
                                            f"residual {float(r['residual']):.1e}", "xlsx_pos_sphere_cpu.py")
                    if fastest:
                        x.font = BOLD
                row += 1
        row += 1
    for ws in (ws_all, ws_tab):
        for j in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(j)].width = 14
    ws_all.auto_filter.ref = ws_all.dimensions
    ws_all.freeze_panes = "A2"
    for j in (12, 13):
        for c in ws_all.iter_rows(min_row=2, min_col=j, max_col=j):
            c[0].number_format = "0.00E+00"

    ws_leg = wb.create_sheet("legend")
    for line in [
        "Tables 5-10 of tex/exact_line_search.tex (tab_cpu_rows_*.tex), CPU-time studies of staging/pos_sphere_cpu.md.",
        "CPU time [ms]: CPU time of GradientDescent::cycle in milliseconds, median of 3 sequential runs with one "
        "deal.II thread; for runs that did not converge, the CPU time until the stop.",
        "status: converged (residual <= 1e-9), not converged, failed (the retraction became undefined).",
        "reason: L = the limit of 1000 iterations; Z = the line search returned step 0; F = failed.",
        "closed form: the exact step is the closed-form ray step (unit sphere, norm-cap); otherwise the "
        "derivative-only root finder (exact) or Armijo.",
        "fastest positive: the fastest converged version on the positive sphere in this row (bold in 'tables').",
        "tables: grey = not converged (reason and residual in the cell comment), red = failed.",
        "Problem: lumped mass, Neumann boundary, V = |x|^2, beta = 100; Q1, P1 level 6, Q2, P2 level 5.",
    ]:
        ws_leg.append([line])
    ws_leg.column_dimensions["A"].width = 120
    wb.save(out)
    print(f"{out}: {ws_all.max_row - 1} rows in 'all', 6 tables")


if __name__ == "__main__":
    main()
